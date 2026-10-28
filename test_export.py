# -*- coding: utf-8 -*-
"""验证导出格式：整行输出 + 差异列标红（Excel / HTML）。"""
import os
import main
from openpyxl import Workbook, load_workbook

HOT = "00FDE2E1"


def mk(path, rows, header):
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    wb.save(path)


def fill_of(cell):
    f = cell.fill
    if f is None or f.fill_type != "solid":
        return None
    c = f.fgColor.rgb
    return c


def main_test():
    base = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_t")
    os.makedirs(base, exist_ok=True)
    pa, pb = os.path.join(base, "a.xlsx"), os.path.join(base, "b.xlsx")
    header = ["编号", "名称", "数量", "状态"]
    mk(pa, [["GC1", "土建", 100, "完成"],
            ["GC2", "机电", 200, "进行中"],
            ["GC3", "消防", 300, "未开始"],
            ["GC4", "精装", 400, "未开始"]], header)
    mk(pb, [["GC1", "土建", 100, "完成"],        # 相同
            ["GC2", "机电", 250, "已完成"],       # 修改：数量、状态
            ["GC4", "精装", 400, "未开始"],       # 相同
            ["GC5", "景观", 500, "未开始"]], header)  # 新增；GC3 删除

    ha, da = main.load_excel(pa)
    hb, db = main.load_excel(pb)
    res = main.compare(ha, da, hb, db, "key", "编号", False, True)
    s = res["stats"]
    assert s == {"same": 2, "changed": 1, "added": 1, "removed": 1}, s

    # ---------------- Excel ----------------
    out = os.path.join(base, "diff.xlsx")
    main.export_excel(res, out, {"file_a": pa, "sheet_a": "Sheet", "file_b": pb,
                                 "sheet_b": "Sheet", "opts": "按关键字列「编号」"})
    wb = load_workbook(out)
    ws = wb["差异明细"]

    # 表头：行1 字段名(合并)，行2 A/B
    assert ws["A1"].value == "差异类型" and ws["B1"].value == "关键字"
    assert ws["C1"].value == "编号" and ws["E1"].value == "名称", (ws["C1"].value, ws["E1"].value)
    assert ws["C2"].value == "A 基准" and ws["D2"].value == "B 对比"
    # 合并区域：A1:A2、B1:B2、C1:D1
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert "A1:A2" in merged and "B1:B2" in merged and "C1:D1" in merged, merged

    # 只输出差异行：3 行数据（1 修改 + 1 新增 + 1 删除）
    data_rows = [r for r in range(3, ws.max_row + 1) if ws.cell(row=r, column=1).value]
    assert len(data_rows) == 3, data_rows

    # 定位第 1 条数据行（修改：GC2，数量 200->250，状态 进行中->已完成）
    r1 = data_rows[0]
    assert ws.cell(row=r1, column=1).value == "修改"
    assert ws.cell(row=r1, column=2).value == "GC2"
    # 整行都在：编号/名称/数量/状态 的 A、B 两侧都有值
    vals = {(i, side): ws.cell(row=r1, column=3 + i * 2 + side).value
            for i in range(4) for side in (0, 1)}
    assert vals[(0, 0)] == "GC2" and vals[(0, 1)] == "GC2", vals
    assert vals[(1, 0)] == "机电" and vals[(1, 1)] == "机电", vals
    assert vals[(2, 0)] == 200 and vals[(2, 1)] == 250, vals     # 数值保持数字类型
    assert vals[(3, 0)] == "进行中" and vals[(3, 1)] == "已完成", vals
    # 差异列标红：数量(2)、状态(3) 两侧都红；编号、名称不红
    for i in (2, 3):
        for side in (0, 1):
            c = ws.cell(row=r1, column=3 + i * 2 + side)
            assert fill_of(c) == HOT, (i, side, fill_of(c))
    for i in (0, 1):
        for side in (0, 1):
            c = ws.cell(row=r1, column=3 + i * 2 + side)
            assert fill_of(c) is None, (i, side, fill_of(c))

    # 新增行：A 侧空（浅灰占位），B 侧有值且全部标红
    r_add = next(r for r in data_rows if ws.cell(row=r, column=1).value == "新增")
    assert all(ws.cell(row=r_add, column=3 + i * 2).value in (None, "") for i in range(4))
    assert all(fill_of(ws.cell(row=r_add, column=3 + i * 2)) == "00F3F4F7" for i in range(4))
    assert ws.cell(row=r_add, column=4).value == "GC5"
    assert all(fill_of(ws.cell(row=r_add, column=4 + i * 2)) == HOT for i in range(4))

    # 删除行：A 侧有值且标红，B 侧空
    r_del = next(r for r in data_rows if ws.cell(row=r, column=1).value == "删除")
    assert ws.cell(row=r_del, column=3).value == "GC3"
    assert all(fill_of(ws.cell(row=r_del, column=3 + i * 2)) == HOT for i in range(4))
    assert all(ws.cell(row=r_del, column=4 + i * 2).value in (None, "") for i in range(4))

    assert ws.freeze_panes == "C3", ws.freeze_panes
    assert "汇总" in wb.sheetnames

    # ---------------- HTML ----------------
    html = main.build_html(res, pa, "Sheet", pb, "Sheet", "按关键字列「编号」")
    assert html.count("class='hot'") == 2 * 2 + 4 + 4, html.count("class='hot'")
    for name in header:
        assert name in html
    assert "data-st='changed'" in html and "data-st='added'" in html and "data-st='removed'" in html
    assert "copyTable()" in html and "A 基准" in html and "B 对比" in html
    assert pa in html and pb in html
    hp = os.path.join(base, "diff.html")
    with open(hp, "w", encoding="utf-8") as f:
        f.write(html)

    print("EXCEL  rows=%d  header_ok  hot_cells_ok  added/removed_ok  freeze=%s"
          % (len(data_rows), ws.freeze_panes))
    print("HTML   hot=%d  filters_ok  copy_btn_ok" % html.count("class='hot'"))
    print("PASS stats=%s" % s)
    print("files:", out, hp)


if __name__ == "__main__":
    main_test()
