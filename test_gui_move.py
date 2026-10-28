# -*- coding: utf-8 -*-
"""端到端验证界面上的「行写回」：选中行 -> A→B -> 源文件真的被改，且界面刷新。"""
import os
import shutil
import time

import tkinter as tk
from openpyxl import Workbook, load_workbook

import main

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_gm")
MSGS = []


def mk(path, rows, header):
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(header)
    for r in rows:
        ws.append(r)
    wb.save(path)


def pump(root, cond, timeout=20):
    t0 = time.time()
    while time.time() - t0 < timeout:
        root.update()
        if cond():
            return True
        time.sleep(0.02)
    return False


def main_test():
    shutil.rmtree(BASE, ignore_errors=True)
    os.makedirs(BASE)
    pa = os.path.join(BASE, "A.xlsx")
    pb = os.path.join(BASE, "B.xlsx")
    header = ["编号", "名称", "数量"]
    mk(pa, [["K1", "甲", 10], ["K2", "乙", 20]], header)
    mk(pb, [["K1", "甲", 10], ["K2", "乙", 99]], header)

    # 打桩对话框
    main.messagebox.askyesno = lambda *a, **k: True
    main.messagebox.showinfo = lambda t, m: MSGS.append(("info", t, m))
    main.messagebox.showwarning = lambda t, m: MSGS.append(("warn", t, m))
    main.messagebox.showerror = lambda t, m: MSGS.append(("err", t, m))

    root = tk.Tk()
    root.withdraw()
    app = main.App(root)
    try:
        app.file_a.set(pa)
        app.file_b.set(pb)
        app.sheet_a.set("数据")
        app.sheet_b.set("数据")
        app.mode.set("key")
        app.key_col.set("编号")
        app.header_row.set(1)
        app.skip_empty.set(True)

        app._start()
        ok = pump(root, lambda: app.result is not None and app.compare_view.aligned)
        assert ok, "比对超时 / 失败: %s" % MSGS
        s = app.result["stats"]
        assert s == {"same": 1, "changed": 1, "added": 0, "removed": 0}, s

        # 选中 K2（修改行）
        idx = [i for i, e in enumerate(app.result["aligned"]) if e["key"] == "K2"][0]
        app.compare_view.sel = {idx}
        app.compare_view.sel_pane = "L"
        app.compare_view._redraw_sel()
        assert app.btn_ab.winfo_exists(), "工具栏 A→B 按钮不存在"

        # 右键菜单按选中行启用/灰化
        app.compare_view._update_move_states()
        st_ab = app.compare_view.menu.entrycget(app.compare_view.mi_ab, "state")
        st_ba = app.compare_view.menu.entrycget(app.compare_view.mi_ba, "state")
        assert st_ab == "normal" and st_ba == "normal", (st_ab, st_ba)

        app._apply_moves("AB")
        ok = pump(root, lambda: "已写入文件B" in app.status.get())
        assert ok, "写回未完成，status=%r msgs=%s" % (app.status.get(), MSGS)

        wb = load_workbook(pb)
        v = wb["数据"].cell(row=3, column=3).value
        wb.close()
        assert v == 20, "B 未被写入, 数量=%r" % v
        assert not [m for m in MSGS if m[0] == "err"], MSGS

        # 写回后界面应立刻刷新为「无差异」
        s2 = app.result["stats"]
        assert s2 == {"same": 2, "changed": 0, "added": 0, "removed": 0}, s2
        assert app.compare_view.aligned and len(app.compare_view.aligned) == 2

        # 再验证反方向 B→A：把 B 改成 77 后重新比对，再写回 A
        wb = load_workbook(pb)
        wb["数据"].cell(row=3, column=3).value = 77
        wb.save(pb)
        wb.close()
        app._start()
        assert pump(root, lambda: app.result and app.result["stats"]["changed"] == 1
                  and app.compare_view.aligned), "重新比对失败"
        idx = [i for i, e in enumerate(app.result["aligned"]) if e["key"] == "K2"][0]
        app.compare_view.sel = {idx}
        app._apply_moves("BA")
        assert pump(root, lambda: "已写入文件A" in app.status.get()), "B→A 未完成"
        wb = load_workbook(pa)
        v = wb["数据"].cell(row=3, column=3).value
        wb.close()
        assert v == 77, "A 未被写入, 数量=%r" % v

        print("status:", app.status.get()[:120])
        print("GUI MOVE A->B / B->A PASSED")
    finally:
        root.destroy()


if __name__ == "__main__":
    main_test()
