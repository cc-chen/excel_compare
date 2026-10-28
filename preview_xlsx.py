# -*- coding: utf-8 -*-
"""把导出的 xlsx 首个工作表还原成 HTML（含合并单元格、底色、粗体），用于目视检查版式。"""
import sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries


def dump(path, outp):
    wb = load_workbook(path)
    ws = wb[wb.sheetnames[0]]

    merged = {}
    covered = set()
    for rng in ws.merged_cells.ranges:
        mn_c, mn_r, mx_c, mx_r = rng.min_col, rng.min_row, rng.max_col, rng.max_row
        merged[(mn_r, mn_c)] = (mx_c - mn_c + 1, mx_r - mn_r + 1)
        for r in range(mn_r, mx_r + 1):
            for c in range(mn_c, mx_c + 1):
                if (r, c) != (mn_r, mn_c):
                    covered.add((r, c))

    def fill(cell):
        f = cell.fill
        if f is not None and f.fill_type == "solid" and f.fgColor and f.fgColor.rgb:
            rgb = f.fgColor.rgb
            if isinstance(rgb, str) and len(rgb) == 8:
                return "#" + rgb[2:]
        return None

    def font(cell):
        f = cell.font
        if f is None:
            return ""
        s = ""
        if f.color is not None and f.color.rgb and isinstance(f.color.rgb, str) and len(f.color.rgb) == 8:
            s += "color:#%s;" % f.color.rgb[2:]
        if f.bold:
            s += "font-weight:700;"
        return s

    rows_html = []
    for r in range(1, ws.max_row + 1):
        tds = []
        for c in range(1, ws.max_column + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(row=r, column=c)
            span = ""
            cs, rs = merged.get((r, c), (1, 1))
            if cs > 1:
                span += " colspan='%d'" % cs
            if rs > 1:
                span += " rowspan='%d'" % rs
            bg = fill(cell)
            st = ""
            if bg:
                st += "background:%s;" % bg
            st += font(cell)
            v = cell.value
            v = "" if v is None else str(v)
            tds.append("<td%s style='%s'>%s</td>" % (span, st, v))
        rows_html.append("<tr>%s</tr>" % "".join(tds))

    # 列宽
    cols = []
    for c in range(1, ws.max_column + 1):
        w = ws.column_dimensions[get_column_letter(c)].width or 10
        cols.append("<col style='width:%dpx'>" % int(w * 7.2))

    html = """<!DOCTYPE html><html><head><meta charset="utf-8">
<style>
 body{margin:16px;font-family:"Microsoft YaHei",Arial,sans-serif;background:#fff;}
 .t{font-size:12px;color:#888;margin-bottom:8px;}
 table{border-collapse:collapse;font-size:12.5px;font-family:"Microsoft YaHei",Arial,sans-serif;}
 td{border:1px solid #d9d9d9;padding:3px 7px;height:22px;}
</style></head><body>
<div class="t">xlsx 还原预览（%s ｜ 工作表「%s」）</div>
<table>%s<tbody>%s</tbody></table>
</body></html>""" % (path.split("/")[-1], ws.title,
                    "<colgroup>%s</colgroup>" % "".join(cols), "".join(rows_html))
    with open(outp, "w", encoding="utf-8") as f:
        f.write(html)
    print("rendered:", outp, "rows:", ws.max_row, "cols:", ws.max_column)


if __name__ == "__main__":
    dump(sys.argv[1], sys.argv[2])
