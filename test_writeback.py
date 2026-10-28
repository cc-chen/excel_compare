# -*- coding: utf-8 -*-
"""验证「行写回」：按行号写回源文件、追加、删除、备份、样式保留、异常保护。"""
import glob
import os
import shutil

import main
from openpyxl import Workbook, load_workbook
from openpyxl.styles import numbers

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_wb")


def mk(path, rows, header):
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(header)
    for r in rows:
        ws.append(r)
    wb.save(path)


def read_all(path, sheet="数据", hr=1):
    h, d, rn = main.load_excel_ex(path, sheet, hr)
    return h, [(k, r) for k, r in zip(rn, d)]


def check(cond, msg):
    if not cond:
        raise AssertionError(msg)


def main_test():
    shutil.rmtree(BASE, ignore_errors=True)
    os.makedirs(BASE)
    pa = os.path.join(BASE, "A.xlsx")
    pb = os.path.join(BASE, "B.xlsx")
    header = ["编号", "名称", "数量", "状态"]

    mk(pa, [["K1", "甲", 10, "完成"],
            ["K2", "乙", 20, "进行中"],
            ["K3", "丙", 30, "未开始"],     # 仅 A 有
            ["K4", "丁", 40, "未开始"]], header)
    mk(pb, [["K1", "甲", 10, "完成"],
            ["K2", "乙", 20, "进行中"],
            ["K4", "丁", 45, "已完成"],     # 与 A 不同
            ["K5", "戊", 50, "未开始"]], header)   # 仅 B 有

    # 给 B 的 数量 列设个自定义格式，验证写回不会破坏格式
    wb = load_workbook(pb)
    for r in range(2, 6):
        wb["数据"].cell(row=r, column=3).number_format = "0.00"
    wb.save(pb)
    wb.close()

    # ---------- 1. 行号映射 ----------
    h, rows = read_all(pa)
    check([r[0] for r in rows] == [2, 3, 4, 5], "rownums A: %s" % [r[0] for r in rows])
    check(rows[0][1][0] == "K1", "first row: %s" % (rows[0],))

    ha, da, ra = main.load_excel_ex(pa, "数据", 1)
    hb, db, rb = main.load_excel_ex(pb, "数据", 1)
    res = main.compare(ha, da, hb, db, "key", "编号", False, True, ra, rb)
    s = res["stats"]
    check(s == {"same": 2, "changed": 1, "added": 1, "removed": 1}, "stats1: %s" % s)
    by_key = {e["key"]: e for e in res["aligned"]}
    check(by_key["K4"]["ra"] == 5 and by_key["K4"]["rb"] == 4, "K4 rows: %s" % by_key["K4"])
    check(by_key["K3"]["ra"] == 4 and by_key["K3"]["rb"] is None, "K3 rows")
    check(by_key["K5"]["ra"] is None and by_key["K5"]["rb"] == 5, "K5 rows")
    print("1 rownum/compare OK  ra/rb=%s" % {k: (v["ra"], v["rb"]) for k, v in by_key.items()})

    # ---------- 2. A -> B：覆盖 + 追加 ----------
    bak = main.backup_file(pb)
    check(os.path.exists(bak), "backup missing")
    edits = []
    e = by_key["K4"]
    edits.append({"op": "set", "row": e["rb"],
                  "values": {c: e["a"].get(c, "") for c in e["changed"]}})
    e = by_key["K3"]
    edits.append({"op": "insert", "row": None, "values": e["a"]})
    done = main.apply_edits(pb, "数据", 1, edits)
    check(len(done) == 2, "done: %s" % (done,))

    h2, rows2 = read_all(pb)
    got = {rn: vals for rn, vals in rows2}
    check(got[4][0] == "K4", "row4 should be K4, got %s" % (got[4],))
    check(got[4][2] == 40 and got[4][3] == "未开始", "K4 not overwritten: %s" % (got[4],))
    check(got[6][0] == "K3" and got[6][1] == "丙", "K3 not appended: %s" % (got.get(6),))
    # 样式保留：数量列自定义格式仍在
    wbv = load_workbook(pb)
    fmt = wbv["数据"].cell(row=4, column=3).number_format
    wbv.close()
    check(fmt == "0.00", "number_format lost: %r" % fmt)
    print("2 A->B set+insert OK  rows=%d  fmt=%s" % (len(rows2), fmt))

    # 重新比对：K4/K3 应变为相同，只剩 K5 仅 B 有
    ha, da, ra = main.load_excel_ex(pa, "数据", 1)
    hb, db, rb = main.load_excel_ex(pb, "数据", 1)
    res = main.compare(ha, da, hb, db, "key", "编号", False, True, ra, rb)
    check(res["stats"] == {"same": 4, "changed": 0, "added": 1, "removed": 0},
          "stats2: %s" % res["stats"])
    print("3 after A->B stats OK %s" % res["stats"])

    # ---------- 3. B -> A：把 K5 追加进 A ----------
    e5 = [x for x in res["aligned"] if x["key"] == "K5"][0]
    main.apply_edits(pa, "数据", 1, [{"op": "insert", "row": None, "values": e5["b"]}])
    ha, da, ra = main.load_excel_ex(pa, "数据", 1)
    hb, db, rb = main.load_excel_ex(pb, "数据", 1)
    res = main.compare(ha, da, hb, db, "key", "编号", False, True, ra, rb)
    check(res["stats"] == {"same": 5, "changed": 0, "added": 0, "removed": 0},
          "stats3: %s" % res["stats"])
    print("4 B->A insert OK  stats=%s" % res["stats"])

    # ---------- 4. 删除 ----------
    e3 = [x for x in res["aligned"] if x["key"] == "K3"][0]
    main.apply_edits(pa, "数据", 1, [{"op": "delete", "row": e3["ra"]}])
    h2, rows3 = read_all(pa)
    keys = [vals[0] for rn, vals in rows3]
    check("K3" not in keys and len(rows3) == 4, "delete failed: %s" % keys)
    print("5 delete OK  A keys=%s" % keys)

    # ---------- 5. 异常保护 ----------
    before = open(pa, "rb").read()
    try:
        main.apply_edits(pa, "数据", 1, [{"op": "set", "row": 1, "values": {"编号": "X"}}])
        raise AssertionError("should raise WriteError")
    except main.WriteError as ex:
        check("表头" in str(ex), "unexpected msg: %s" % ex)
    after = open(pa, "rb").read()
    check(before == after, "file changed despite error!")

    try:
        main.apply_edits(os.path.join(BASE, "x.xls"), "数据", 1, [])
        raise AssertionError("should reject .xls")
    except main.WriteError as ex:
        check(".xls" in str(ex), "unexpected msg: %s" % ex)
    print("6 guard OK  表头越界不改文件 / .xls 拒写")

    # ---------- 6. 备份文件 ----------
    baks = glob.glob(os.path.join(BASE, "B.bak-*"))
    check(len(baks) == 1, "bak files: %s" % baks)
    wb = load_workbook(baks[0])
    orig = [c.value for c in wb["数据"][4]]
    wb.close()
    check(orig[2] == 45, "backup should hold original 45, got %s" % orig[2])
    print("7 backup OK  %s (原始 K4 数量=%s)" % (os.path.basename(baks[0]), orig[2]))

    # ---------- 7. 按行号模式(index)的写回 ----------
    pi = os.path.join(BASE, "I1.xlsx")
    pj = os.path.join(BASE, "I2.xlsx")
    mk(pi, [["a", 1], ["b", 2], ["c", 3]], ["列1", "列2"])
    mk(pj, [["a", 1], ["b", 9], ["c", 3]], ["列1", "列2"])
    hi, di, ri = main.load_excel_ex(pi, "数据", 1)
    hj, dj, rj = main.load_excel_ex(pj, "数据", 1)
    r = main.compare(hi, di, hj, dj, "index", "", False, True, ri, rj)
    check(r["stats"]["changed"] == 1, "index stats: %s" % r["stats"])
    ch = [x for x in r["aligned"] if x["status"] == "changed"][0]
    check(ch["ra"] == 3 and ch["rb"] == 3, "index rows: %s" % ((ch["ra"], ch["rb"]),))
    main.apply_edits(pj, "数据", 1, [{"op": "set", "row": ch["rb"],
                                    "values": {c: ch["a"].get(c) for c in ch["changed"]}}])
    hi, di, ri = main.load_excel_ex(pi, "数据", 1)
    hj, dj, rj = main.load_excel_ex(pj, "数据", 1)
    r = main.compare(hi, di, hj, dj, "index", "", False, True, ri, rj)
    check(r["stats"]["changed"] == 0, "index after write: %s" % r["stats"])
    print("8 index-mode write OK  stats=%s" % r["stats"])

    print("WRITEBACK ALL PASSED")


if __name__ == "__main__":
    main_test()
