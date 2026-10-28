# -*- coding: utf-8 -*-
"""
ExcelDiff - 两个 Excel 内容对比工具（类 Beyond Compare 双栏视图）
支持 .xlsx / .xlsm / .xls
比对方式：按行号对齐 或 按关键字列
左栏 = 文件A(基准)，右栏 = 文件B(对比)；差异项红色高亮，中间状态栏显示 ≠/</>
可导出差异 Excel 与 HTML 报告
"""
import os
import sys
import json
import queue
import shutil
import datetime
import threading
import webbrowser
import tkinter as tk
from copy import copy
from tkinter import ttk, filedialog, messagebox

from openpyxl import load_workbook
import xlrd

# ----------------------------------------------------------------------------
# 比对引擎（纯函数，可独立测试）
# ----------------------------------------------------------------------------

def norm_val(v, ignore_case=False):
    """把单元格值规范成可比较的字符串。"""
    if v is None:
        return ""
    if isinstance(v, float):
        if v == int(v):
            v = int(v)
    s = str(v)
    if ignore_case:
        s = s.lower()
    return s.strip()


def load_excel(path, sheet_name=None, header_row=1):
    """读取 Excel，返回 (headers, data_rows)。"""
    h, d, _ = load_excel_ex(path, sheet_name, header_row)
    return h, d


def load_excel_ex(path, sheet_name=None, header_row=1):
    """读取 Excel，返回 (headers, data_rows, rownums)。

    rownums[i] 是 data_rows[i] 在源工作表中的真实行号（1-based），
    用于把界面上的行写回原文件。整行为空的过渡行 openpyxl 会给出空元组，
    此时按上一行 +1 递推，保证行号与源文件一致。
    """
    ext = os.path.splitext(path)[1].lower()
    rownums = []
    if ext in (".xlsx", ".xlsm"):
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet_name:
                ws = wb[sheet_name]
            else:
                ws = wb.active
            # 关键：read_only 模式会信任工作表里的 <dimension> 声明；
            # 很多工具(腾讯文档/WPS/导出工具)导出的文件 dimension 是错的(如 "A1")，
            # 会导致只读到 1 行。重置尺寸后按实际 XML 内容流式读取全部行。
            if hasattr(ws, "reset_dimensions"):
                try:
                    ws.reset_dimensions()
                except Exception:
                    pass
            rows = []
            for row in ws.iter_rows():
                if row:
                    rn = getattr(row[0], "row", None)
                    if not isinstance(rn, int):
                        rn = (rownums[-1] + 1) if rownums else 1
                    rows.append(tuple(c.value for c in row))
                else:                      # 整行为空
                    rn = (rownums[-1] + 1) if rownums else 1
                    rows.append(())
                rownums.append(rn)
        finally:
            wb.close()
    elif ext == ".xls":
        book = xlrd.open_workbook(path)
        if sheet_name:
            sh = book.sheet_by_name(sheet_name)
        else:
            sh = book.sheet_by_index(0)
        rows = []
        for r in range(sh.nrows):
            rows.append([sh.cell_value(r, c) for c in range(sh.ncols)])
            rownums.append(r + 1)
    else:
        raise ValueError("不支持的文件格式：%s（仅支持 .xlsx/.xlsm/.xls）" % ext)

    hr = max(1, int(header_row))
    if rows and len(rows) >= hr:
        raw_header = rows[hr - 1]
        headers = [("" if v is None else str(v)) for v in raw_header]
        data = rows[hr:]
        rnums = rownums[hr:]
    else:
        headers = []
        data = rows
        rnums = list(rownums)

    ncol = max(len(headers), max((len(r) for r in data), default=0))
    norm = []
    for r in data:
        row = [("" if v is None else v) for v in r]
        if len(row) < ncol:
            row += [""] * (ncol - len(row))
        norm.append(row)
    if len(headers) < ncol:
        headers += [""] * (ncol - len(headers))
    if len(rnums) < len(norm):          # 兜底，理论上不会发生
        base = rnums[-1] if rnums else hr
        rnums = rnums + [base + i + 1 for i in range(len(norm) - len(rnums))]
    return headers, norm, rnums[:len(norm)]


def _build_index_rows(data, unified, rownums=None):
    rows = []
    for i, r in enumerate(data):
        d = {lab: (r[j] if j < len(r) else "") for j, lab in enumerate(unified)}
        rn = rownums[i] if (rownums and i < len(rownums)) else None
        rows.append((("#", i), d, rn))
    return rows


def _build_key_rows(data, unified, headers, key_col, rownums=None):
    posmap = {h: idx for idx, h in enumerate(headers)}
    rows = []
    for i, r in enumerate(data):
        d = {}
        for h in unified:
            p = posmap.get(h)
            d[h] = (r[p] if (p is not None and p < len(r)) else "")
        k = str(d.get(key_col, ""))
        rn = rownums[i] if (rownums and i < len(rownums)) else None
        rows.append((k, d, rn))
    return rows


def compare(headers_a, data_a, headers_b, data_b,
            mode, key_col, ignore_case=False, skip_empty=False,
            rownums_a=None, rownums_b=None):
    """
    返回 dict：
    {
      "unified": [列名...],
      "diffs":  [ {status, key, col, a, b}... ],   # 供导出/旧列表使用
      "aligned":[ {status, key, a, b, changed:[列], ra, rb}... ],
                # ra/rb：该行在文件A/文件B 工作表中的真实行号（用于写回）
      "stats":  {same, changed, added, removed},
      "warnings": [str...],
      "struct": {only_a:[], only_b:[]}
    }
    """
    warnings = []

    if mode == "index":
        ncol = max(len(headers_a),
                   max((len(r) for r in data_a), default=0),
                   len(headers_b),
                   max((len(r) for r in data_b), default=0))
        unified = []
        for i in range(ncol):
            lab = headers_a[i] if i < len(headers_a) and headers_a[i] != "" else "列%d" % (i + 1)
            unified.append(lab)
        rows_a = _build_index_rows(data_a, unified, rownums_a)
        rows_b = _build_index_rows(data_b, unified, rownums_b)
        shared = list(unified)
        struct = {"only_a": [], "only_b": []}
        if len(headers_a) != len(headers_b):
            warnings.append("两个文件的列数不同（A=%d, B=%d），按行号对齐比较。"
                            % (len(headers_a), len(headers_b)))
    else:
        set_a = [h for h in headers_a if h != ""]
        set_b = [h for h in headers_b if h != ""]
        unified = list(dict.fromkeys(set_a + [h for h in set_b if h not in set_a]))
        rows_a = _build_key_rows(data_a, unified, headers_a, key_col, rownums_a)
        rows_b = _build_key_rows(data_b, unified, headers_b, key_col, rownums_b)
        shared = [c for c in unified if (c in set_a and c in set_b)]
        only_a = [c for c in set_a if c not in set_b]
        only_b = [c for c in set_b if c not in set_a]
        struct = {"only_a": only_a, "only_b": only_b}
        if only_a:
            warnings.append("以下列仅存在于文件A：%s" % "、".join(only_a))
        if only_b:
            warnings.append("以下列仅存在于文件B：%s" % "、".join(only_b))
        dup_a = [k for k, _, _ in rows_a if [x for x, _, _ in rows_a].count(k) > 1]
        dup_b = [k for k, _, _ in rows_b if [x for x, _, _ in rows_b].count(k) > 1]
        if dup_a:
            warnings.append("文件A存在重复关键字：%s" % "、".join(dict.fromkeys(dup_a)))
        if dup_b:
            warnings.append("文件B存在重复关键字：%s" % "、".join(dict.fromkeys(dup_b)))

    map_a = {k: (d, rn) for k, d, rn in rows_a}
    map_b = {k: (d, rn) for k, d, rn in rows_b}
    keys = list(dict.fromkeys([k for k, _, _ in rows_a] + [k for k, _, _ in rows_b]))

    def is_skip(d):
        if not skip_empty:
            return False
        return all(norm_val(v) == "" for v in d.values())

    ic = ignore_case
    aligned = []
    for k in keys:
        in_a, in_b = k in map_a, k in map_b
        if in_a and in_b:
            (da, ra), (db, rb) = map_a[k], map_b[k]
            if is_skip(da) and is_skip(db):
                continue
            changed = [c for c in shared
                       if norm_val(da.get(c, ""), ic) != norm_val(db.get(c, ""), ic)]
            if changed:
                aligned.append({"status": "changed", "key": _key_label(k),
                                "a": da, "b": db, "changed": changed,
                                "ra": ra, "rb": rb})
            else:
                aligned.append({"status": "same", "key": _key_label(k),
                                "a": da, "b": db, "changed": [],
                                "ra": ra, "rb": rb})
        elif in_a:
            if is_skip(map_a[k][0]):
                continue
            aligned.append({"status": "removed", "key": _key_label(k),
                            "a": map_a[k][0], "b": None, "changed": [],
                            "ra": map_a[k][1], "rb": None})
        else:
            if is_skip(map_b[k][0]):
                continue
            aligned.append({"status": "added", "key": _key_label(k),
                            "a": None, "b": map_b[k][0], "changed": [],
                            "ra": None, "rb": map_b[k][1]})

    # 由 aligned 派生 diffs / stats（保持旧接口形态）
    stats = {"same": 0, "changed": 0, "added": 0, "removed": 0}
    diffs = []
    for e in aligned:
        st = e["status"]
        stats[st] += 1
        if st == "changed":
            for c in e["changed"]:
                diffs.append({"status": "修改", "key": e["key"], "col": c,
                              "a": e["a"].get(c, ""), "b": e["b"].get(c, "")})
        elif st == "removed":
            for c in unified:
                diffs.append({"status": "删除", "key": e["key"], "col": c,
                              "a": e["a"].get(c, ""), "b": ""})
        elif st == "added":
            for c in unified:
                diffs.append({"status": "新增", "key": e["key"], "col": c,
                              "a": "", "b": e["b"].get(c, "")})

    return {"unified": unified, "diffs": diffs, "aligned": aligned,
            "stats": stats, "warnings": warnings, "struct": struct}


def _key_label(k):
    return k[1] if isinstance(k, tuple) else k


# ----------------------------------------------------------------------------
# 写回源文件（界面上把行"移到"另一侧）
# ----------------------------------------------------------------------------

class WriteError(Exception):
    pass


def backup_file(path):
    """把文件复制一份 <名>.bak-时间戳<扩展名>，返回备份路径。

    保留原扩展名，这样备份文件可以直接双击用 Excel 打开。
    """
    stem, ext = os.path.splitext(path)
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    bak = "%s.bak-%s%s" % (stem, stamp, ext)
    i = 1
    while os.path.exists(bak):
        bak = "%s.bak-%s-%d%s" % (stem, stamp, i, ext)
        i += 1
    shutil.copy2(path, bak)
    return bak


def _write_val(v):
    if v is None:
        return None
    if isinstance(v, str) and v.strip() == "":
        return None
    return v


def _sheet_colmap(ws, header_row):
    """表头名 -> 列号(1-based)。"""
    cmap = {}
    for c in range(1, (ws.max_column or 1) + 1):
        v = ws.cell(row=header_row, column=c).value
        if v is None:
            continue
        k = str(v).strip()
        if k and k not in cmap:
            cmap[k] = c
    return cmap


def _fill_row(ws, row, values, cmap):
    n = 0
    for col, v in values.items():
        c = cmap.get(str(col))
        if not c:
            continue
        ws.cell(row=row, column=c).value = _write_val(v)
        n += 1
    return n


def _clone_row_style(ws, src_row, dst_row, last_col):
    for c in range(1, last_col + 1):
        src = ws.cell(row=src_row, column=c)
        if src.has_style:
            ws.cell(row=dst_row, column=c)._style = copy(src._style)


def apply_edits(path, sheet, header_row, edits):
    """把编辑批量写入 Excel 并保存。

    edits: [{"op": "set"|"insert"|"delete", "row": 行号(1-based),
             "values": {列名: 值}}]
    执行顺序：追加 -> 覆盖(用原行号) -> 删除(行号降序)，三者互不干扰。
    返回 [(op, 行号, 写入列数), ...]

    注意：任何异常都在 wb.save() 之前抛出，此时文件尚未被改动。
    """
    ext = os.path.splitext(path)[1].lower()
    if ext not in (".xlsx", ".xlsm"):
        raise WriteError("不支持写回 %s 格式（只有 .xlsx / .xlsm 能写）。"
                         "请先用 Excel 另存为 .xlsx。" % (ext or "该"))
    wb = load_workbook(path, keep_vba=(ext == ".xlsm"))
    try:
        ws = wb[sheet] if (sheet and sheet in wb.sheetnames) else wb.active
        cmap = _sheet_colmap(ws, header_row)
        if not cmap:
            raise WriteError("目标工作表第 %d 行没有表头，无法定位要写哪一列。" % header_row)
        last_col = max(cmap.values())

        appends = [e for e in edits if e["op"] == "insert"]
        updates = [e for e in edits if e["op"] == "set"]
        deletes = sorted([e for e in edits if e["op"] == "delete"],
                         key=lambda e: int(e["row"]), reverse=True)
        done = []

        for e in appends:
            row = ws.max_row + 1
            if row - 1 > header_row:
                _clone_row_style(ws, row - 1, row, last_col)
            done.append(("insert", row, _fill_row(ws, row, e["values"], cmap)))

        for e in updates:
            row = int(e["row"])
            if row <= header_row:
                raise WriteError("行号 %d 落在表头区域内，已中止（文件未被改动）。" % row)
            done.append(("set", row, _fill_row(ws, row, e["values"], cmap)))

        for e in deletes:
            row = int(e["row"])
            if row <= header_row:
                raise WriteError("行号 %d 落在表头区域内，已中止（文件未被改动）。" % row)
            ws.delete_rows(row, 1)
            done.append(("delete", row, 0))

        wb.save(path)
        return done
    finally:
        wb.close()


# ----------------------------------------------------------------------------
# 导出：差异 Excel
# ----------------------------------------------------------------------------

def _to_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v)


STATUS_LABEL = {"same": "相同", "changed": "修改", "added": "新增", "removed": "删除"}
# 状态单元格配色：底色, 字色
STATUS_STYLE = {
    "修改": ("FFF2CC", "8A5A00"),
    "新增": ("E2EFDA", "2E7D32"),
    "删除": ("FCE4E4", "C62828"),
}
HOT_FILL = "FDE2E1"   # 差异单元格底色（浅红）
HOT_FONT = "C00000"   # 差异单元格字色（深红）


def _diff_rows(result):
    """取出需要导出的行：完整行 + 差异列，跳过完全相同的行。"""
    return [e for e in result["aligned"] if e["status"] != "same"]


def _hot_cols(entry):
    """该行需要标红的列集合。"""
    if entry["status"] == "changed":
        return set(entry["changed"])
    return set(entry["a"] if entry["a"] is not None else entry["b"])


def export_excel(result, path, info=None):
    """导出差异 Excel：每行一条完整记录，A/B 值并排，差异单元格标红。

    版式：
      行1  差异类型 | 关键字 | 字段1(跨2列) | 字段2(跨2列) | ...
      行2  差异类型 | 关键字 |  A 基准 | B 对比 | A | B | ...
      行3+ 数据
    """
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    info = info or {}
    unified = result["unified"]
    rows = _diff_rows(result)
    stats = result["stats"]

    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_font = Font(bold=True, color="555555")
    hdr_fill = PatternFill("solid", fgColor="F2F4F7")
    ab_fill = PatternFill("solid", fgColor="FAFBFC")
    red_font = Font(color=HOT_FONT, bold=True)
    red_fill = PatternFill("solid", fgColor=HOT_FILL)
    na_fill = PatternFill("solid", fgColor="F3F4F7")   # 该侧不存在的单元格
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center", wrap_text=False)

    wb = Workbook()
    ws = wb.active
    ws.title = "差异明细"

    def col_of(i, side):
        """第 i 个字段的 A/B 值所在列（side: 0=A, 1=B）。"""
        return 3 + i * 2 + side

    last_col = 2 + len(unified) * 2

    # ---- 表头（先建格子再合并，避免合并后无法逐格设样式）----
    ws.cell(row=1, column=1, value="差异类型")
    ws.cell(row=1, column=2, value="关键字")
    ws.cell(row=2, column=1, value=None)
    ws.cell(row=2, column=2, value=None)
    for i, name in enumerate(unified):
        c0 = col_of(i, 0)
        ws.cell(row=1, column=c0, value=str(name))
        ws.cell(row=2, column=c0, value="A 基准")
        ws.cell(row=2, column=c0 + 1, value="B 对比")
    for r, fill in ((1, hdr_fill), (2, ab_fill)):
        for c in range(1, last_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = hdr_font
            cell.fill = fill
            cell.alignment = center
            cell.border = border
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
    for i in range(len(unified)):
        c0 = col_of(i, 0)
        ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + 1)

    # ---- 数据行：整行输出，差异列标红 ----
    r = 3
    for e in rows:
        st = e["status"]
        hot = _hot_cols(e)
        label = STATUS_LABEL[st]
        bg, fg = STATUS_STYLE.get(label, ("FFFFFF", "333333"))

        c_status = ws.cell(row=r, column=1, value=label)
        c_status.fill = PatternFill("solid", fgColor=bg)
        c_status.font = Font(bold=True, color=fg)
        c_status.alignment = center
        c_status.border = border

        c_key = ws.cell(row=r, column=2, value=e["key"])
        c_key.font = Font(bold=True, color="333333")
        c_key.alignment = left
        c_key.border = border

        for i, name in enumerate(unified):
            va = "" if e["a"] is None else e["a"].get(name, None)
            vb = "" if e["b"] is None else e["b"].get(name, None)
            for side, val, missing in ((0, va, e["a"] is None), (1, vb, e["b"] is None)):
                cell = ws.cell(row=r, column=col_of(i, side))
                if not missing:
                    if isinstance(val, (int, float)) and not isinstance(val, bool):
                        cell.value = val
                    else:
                        cell.value = _to_str(val)
                cell.alignment = left
                cell.border = border
                if missing:
                    cell.fill = na_fill
                elif name in hot:
                    cell.fill = red_fill
                    cell.font = red_font
        r += 1

    # 列宽 / 冻结 / 视图
    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 16
    for i in range(len(unified)):
        for side in (0, 1):
            ws.column_dimensions[ws.cell(row=2, column=col_of(i, side))
                                .column_letter].width = 16
    ws.freeze_panes = "C3"
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 18
    ws.sheet_view.zoomScale = 90
    if len(rows) == 0:
        ws.cell(row=3, column=1, value="（没有差异，两个文件的指定工作表内容一致）")

    # ---- 汇总页 ----
    ws2 = wb.create_sheet("汇总")
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 60
    ws2.append(["项目", "内容"])
    meta = [
        ("文件A（基准）", "%s  [%s]" % (info.get("file_a", ""), info.get("sheet_a", ""))),
        ("文件B（对比）", "%s  [%s]" % (info.get("file_b", ""), info.get("sheet_b", ""))),
        ("比对方式", info.get("opts", "")),
        ("生成时间", info.get("now", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"))),
        ("", ""),
        ("相同行", stats["same"]),
        ("修改行", stats["changed"]),
        ("新增行", stats["added"]),
        ("删除行", stats["removed"]),
        ("差异行合计", stats["changed"] + stats["added"] + stats["removed"]),
    ]
    for k, v in meta:
        ws2.append([k, v])
    for cell in ws2["A"]:
        cell.font = Font(bold=True)
    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    wb.save(path)


# ----------------------------------------------------------------------------
# 导出：HTML 报告
# ----------------------------------------------------------------------------

def _esc(s):
    return (str(s).replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;").replace("\n", "<br>"))


CSS = """
 body{font-family:-apple-system,"Segoe UI","Microsoft YaHei",sans-serif;margin:0;padding:20px 24px 40px;color:#222;background:#fff;}
 h1{font-size:19px;margin:0 0 4px;}
 h2{font-size:14px;margin:20px 0 8px;color:#333;}
 .k{color:#888;font-size:12px;}
 .box{background:#fafbfc;border:1px solid #e8eaee;border-radius:8px;padding:12px 16px;margin:10px 0;font-size:13px;line-height:1.9;}
 .chips span{display:inline-block;font-size:13px;padding:3px 12px;border-radius:14px;margin-right:8px;font-weight:600;}
 .c-same{background:#eef1f5;color:#44515f;}
 .c-mod{background:#fff2cc;color:#8a5a00;}
 .c-add{background:#e2efda;color:#2e7d32;}
 .c-del{background:#fce4e4;color:#c62828;}
 .bar{display:flex;align-items:center;gap:8px;margin:14px 0 8px;flex-wrap:wrap;}
 .bar button{border:1px solid #d7dbe2;background:#fff;border-radius:6px;padding:5px 12px;font-size:12.5px;cursor:pointer;color:#333;font-family:inherit;}
 .bar button.on{background:#2f6fdf;border-color:#2f6fdf;color:#fff;}
 .bar .sp{flex:1;}
 .bar .hint{font-size:12px;color:#2e7d32;min-width:180px;}
 .wrap{display:inline-block;max-width:100%;overflow:auto;max-height:calc(100vh - 120px);border:1px solid #e3e6ec;border-radius:8px;}
 table{border-collapse:separate;border-spacing:0;font-size:12.5px;white-space:nowrap;}
 th,td{border-right:1px solid #edf0f4;border-bottom:1px solid #edf0f4;padding:5px 9px;text-align:left;}
 thead{position:sticky;top:0;z-index:4;}
 thead th{background:#f2f4f7;color:#555;font-weight:700;text-align:center;font-size:12px;line-height:1.5;}
 thead tr:nth-child(2) th{background:#f7f9fb;font-weight:600;color:#8a94a6;font-size:11.5px;border-bottom:1px solid #dfe3ea;}
 th.ab-a{color:#8a94a6;}
 th.ab-b{color:#5b6b85;}
 td{background:#fff;vertical-align:middle;}
 td.hot{background:#fde2e1;color:#c00000;font-weight:700;}
 tbody tr:hover td{background:#f5f9ff;}
 tbody tr:hover td.hot{background:#fbd7d5;}
 td.st{font-weight:700;text-align:center;}
 td.st-mod{background:#fff2cc;color:#8a5a00;}
 td.st-add{background:#e2efda;color:#2e7d32;}
 td.st-del{background:#fce4e4;color:#c62828;}
 td.key{font-weight:700;color:#2a3240;}
 td.na{background:#f3f4f7;}
 .fix0{position:sticky;left:0;z-index:2;width:52px;min-width:52px;}
 .fix1{position:sticky;left:52px;z-index:2;}
 thead .fix0,thead .fix1{z-index:6;}
 td.fix1{background:#fbfcfd;}
 tbody tr:hover td.fix1{background:#f5f9ff;}
 .nada{padding:14px;color:#888;font-size:13px;}
 .lgs{font-size:12px;color:#666;margin-top:10px;}
 .lg{display:inline-block;padding:1px 8px;border-radius:3px;margin-right:10px;}
"""

JS = """
function flt(btn, f){
  document.querySelectorAll('.bar button[data-f]').forEach(function(b){b.classList.toggle('on', b===btn);});
  document.querySelectorAll('#tbl tbody tr').forEach(function(tr){
    tr.style.display = (f==='all' || tr.dataset.st===f) ? '' : 'none';
  });
}
function flash(msg){
  var h=document.getElementById('hint'); h.textContent=msg;
  setTimeout(function(){h.textContent='';}, 2600);
}
function copyTable(){
  var lines=[];
  // 表头由报告生成时直接给出（字段名 + A/B 侧），避免解析 rowspan/colspan 出错
  var head=(window.HDR||[]).map(function(p){
    return p[1] ? (p[0] + '\\uff08' + p[1] + '\\uff09') : p[0];
  });
  if(head.length){ lines.push(head.join('\\t')); }
  document.querySelectorAll('#tbl tbody tr').forEach(function(tr){
    if(tr.style.display==='none') return;
    var cells=[];
    tr.querySelectorAll('td').forEach(function(td){ cells.push((td.textContent||'').replace(/\\r?\\n/g,' ')); });
    lines.push(cells.join('\\t'));
  });
  var text=lines.join('\\r\\n');
  if(navigator.clipboard && navigator.clipboard.writeText){
    navigator.clipboard.writeText(text).then(function(){
      flash('已复制 '+lines.length+' 行，可直接粘贴到 Excel');
    }, function(){ fallback(text); });
  } else { fallback(text); }
}
function fallback(text){
  var ta=document.createElement('textarea');
  ta.value=text; document.body.appendChild(ta); ta.select();
  try{ document.execCommand('copy'); flash('已复制，可直接粘贴到 Excel'); }
  catch(e){ flash('复制失败，请手动选择表格复制'); }
  document.body.removeChild(ta);
}
"""


def build_html(result, file_a, sheet_a, file_b, sheet_b, opts):
    """HTML 报告：每行一条完整记录，A/B 值并排，差异单元格标红。"""
    stats = result["stats"]
    unified = result["unified"]
    rows = _diff_rows(result)
    warnings = result["warnings"]
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    total_diff = stats["changed"] + stats["added"] + stats["removed"]

    # ---- 表头 ----
    head1 = ["<th class='fix0 st' rowspan='2'>差异类型</th>",
             "<th class='fix1' rowspan='2'>关键字</th>"]
    head1 += ["<th colspan='2'>%s</th>" % _esc(n) for n in unified]
    head2 = []
    for _ in unified:
        head2.append("<th class='ab-a'>A 基准</th><th class='ab-b'>B 对比</th>")
    thead = ("<thead><tr>%s</tr><tr>%s</tr></thead>"
             % ("".join(head1), "".join(head2)))

    # ---- 数据行：整行 + 差异列标红 ----
    body = []
    for e in rows:
        st = e["status"]
        hot = _hot_cols(e)
        label = STATUS_LABEL[st]
        cls = {"changed": "st-mod", "added": "st-add", "removed": "st-del"}[st]
        tds = ["<td class='fix0 st %s'>%s</td>" % (cls, label),
               "<td class='fix1 key'>%s</td>" % _esc(e["key"])]
        for name in unified:
            for side, missing in ((0, e["a"] is None), (1, e["b"] is None)):
                val = "" if missing else _to_str((e["a"] if side == 0 else e["b"]).get(name, ""))
                if missing:
                    tds.append("<td class='na'></td>")
                elif name in hot:
                    tds.append("<td class='hot'>%s</td>" % _esc(val))
                else:
                    tds.append("<td>%s</td>" % _esc(val))
        body.append("<tr data-st='%s'>%s</tr>" % (st, "".join(tds)))

    tbody = ("<tbody>%s</tbody>" % "".join(body)) if body else \
            "<tbody><tr><td class='nada' colspan='%d'>没有差异，两个文件的指定工作表内容一致。</td></tr></tbody>" \
            % (2 + len(unified) * 2)

    warn_html = "".join("<li>%s</li>" % _esc(w) for w in warnings) or "<li>无</li>"

    # 复制表格时的表头（字段名, A/B 侧），由程序直接给出，与表格结构严格对应
    hdr_pairs = [["差异类型", ""], ["关键字", ""]]
    for n in unified:
        hdr_pairs.append([str(n), "A 基准"])
        hdr_pairs.append([str(n), "B 对比"])
    hdr_js = json.dumps(hdr_pairs, ensure_ascii=False)

    return """<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<title>Excel 对比报告</title>
<style>%s</style></head><body>
<h1>Excel 对比报告</h1>
<div class="k">生成时间：%s</div>

<div class="box">
 <div><b>文件A（基准）：</b>%s &nbsp;<span class="k">[%s]</span></div>
 <div><b>文件B（对比）：</b>%s &nbsp;<span class="k">[%s]</span></div>
 <div><b>比对方式：</b>%s</div>
</div>

<div class="box chips">
 <span class="c-same">相同 %d</span>
 <span class="c-mod">修改 %d</span>
 <span class="c-add">新增 %d</span>
 <span class="c-del">删除 %d</span>
 <span class="c-same">差异行合计 %d</span>
</div>

<h2>结构 / 注意事项</h2>
<ul>%s</ul>

<h2>差异明细（只列差异行，整行输出；红色为差异单元格）</h2>
<div class="bar">
 <button class="on" data-f="all" onclick="flt(this,'all')">全部 %d</button>
 <button data-f="changed" onclick="flt(this,'changed')">修改 %d</button>
 <button data-f="added" onclick="flt(this,'added')">新增 %d</button>
 <button data-f="removed" onclick="flt(this,'removed')">删除 %d</button>
 <span class="sp"></span>
 <span id="hint" class="hint"></span>
 <button onclick="copyTable()">复制表格（粘贴到 Excel）</button>
</div>
<div class="wrap"><table id="tbl">%s%s</table></div>
<div class="lgs">
 <span class="lg" style="background:#fde2e1;color:#c00000;font-weight:700">差异单元格</span>
 <span class="lg" style="background:#fff2cc;color:#8a5a00">修改</span>
 <span class="lg" style="background:#e2efda;color:#2e7d32">仅 B 有</span>
 <span class="lg" style="background:#fce4e4;color:#c62828">仅 A 有</span>
 &nbsp;·&nbsp;浅灰格表示该侧没有这条记录；滚动时左侧「差异类型 / 关键字」固定；点上方按钮可只看某一类差异。
</div>
<script>var HDR=%s;</script>
<script>%s</script>
</body></html>""" % (
        CSS, now, _esc(file_a), _esc(sheet_a), _esc(file_b), _esc(sheet_b),
        _esc(opts), stats["same"], stats["changed"], stats["added"], stats["removed"],
        total_diff, warn_html, len(rows),
        stats["changed"], stats["added"], stats["removed"],
        thead, tbody, hdr_js, JS)


# ----------------------------------------------------------------------------
# 双栏对比视图（类 Beyond Compare）
# ----------------------------------------------------------------------------

class CompareView(tk.Frame):
    """左文件A / 右文件B 双栏，差异项红色高亮，中间状态栏显示 ≠/</>。"""

    def __init__(self, master):
        super().__init__(master)
        self.aligned = []
        self.cols = []
        self.colw = []
        self.idx_w = 48
        self.rowh = 22
        self.hdrh = 26
        self.gw = 30
        self.font = ("Microsoft YaHei", 9)
        self.total_w = 0
        # 选中状态
        self.sel = set()          # 选中的 aligned 行下标
        self.sel_anchor = None    # 范围选择锚点
        self.sel_pane = "L"       # 最近一次点击所在栏: L / R
        self._dragging = False
        self.on_status = None     # 状态栏回调（由 App 注入）
        self.on_move = None       # 行移动回调(direction, indices)，由 App 注入

        self.lblL = ttk.Label(self, text="文件A（基准）", font=("Microsoft YaHei", 9, "bold"))
        self.lblR = ttk.Label(self, text="点击选行 · Ctrl/Shift 多选 · Ctrl+C 复制 · Ctrl+→ A写B · Ctrl+← B写A ｜ ≠修改 <仅A >仅B",
                              font=("Microsoft YaHei", 9), anchor="e")

        self.hdrL = tk.Canvas(self, height=self.hdrh, bg="#eaeef3", highlightthickness=0)
        self.hdrR = tk.Canvas(self, height=self.hdrh, bg="#eaeef3", highlightthickness=0)
        self.gHdr = tk.Canvas(self, width=self.gw, height=self.hdrh, bg="#eaeef3", highlightthickness=0)

        self.cvL = tk.Canvas(self, bg="white", highlightthickness=0)
        self.cvR = tk.Canvas(self, bg="white", highlightthickness=0)
        self.gutter = tk.Canvas(self, width=self.gw, bg="#f6f6f6", highlightthickness=0)

        self.vsb = ttk.Scrollbar(self, orient="vertical", command=self._vscroll)
        self.hsbL = ttk.Scrollbar(self, orient="horizontal", command=self._hscrollL)
        self.hsbR = ttk.Scrollbar(self, orient="horizontal", command=self._hscrollR)

        self._layout()
        self._bind_scroll()
        self._build_menu()

    def _layout(self):
        self.rowconfigure(2, weight=1)
        self.columnconfigure(0, weight=1)
        self.columnconfigure(2, weight=1)
        self.lblL.grid(row=0, column=0, sticky="w", padx=4, pady=(2, 0))
        self.lblR.grid(row=0, column=2, sticky="e", padx=4, pady=(2, 0))
        self.hdrL.grid(row=1, column=0, sticky="ew")
        self.gHdr.grid(row=1, column=1, sticky="ns")
        self.hdrR.grid(row=1, column=2, sticky="ew")
        self.cvL.grid(row=2, column=0, sticky="nsew")
        self.gutter.grid(row=2, column=1, sticky="ns")
        self.cvR.grid(row=2, column=2, sticky="nsew")
        self.hsbL.grid(row=3, column=0, sticky="ew")
        self.hsbR.grid(row=3, column=2, sticky="ew")
        self.vsb.grid(row=2, column=3, sticky="ns")

    def _bind_scroll(self):
        for cv in (self.cvL, self.cvR, self.gutter):
            cv.bind("<MouseWheel>", lambda e, c=cv: self._wheel(e, c))
            cv.bind("<Shift-MouseWheel>", lambda e, c=cv: self._hwheel(e, c))
        # 选中行 / 复制相关事件
        for cv, pane in ((self.cvL, "L"), (self.cvR, "R")):
            cv.configure(takefocus=1)
            cv.bind("<Button-1>", lambda e, c=cv, p=pane: self._on_press(e, c, p))
            cv.bind("<B1-Motion>", lambda e, c=cv: self._on_motion(e, c))
            cv.bind("<ButtonRelease-1>", lambda e, c=cv: self._on_release(e, c))
            cv.bind("<Control-Button-1>", lambda e, c=cv, p=pane: self._on_ctrl(e, c, p))
            cv.bind("<Shift-Button-1>", lambda e, c=cv, p=pane: self._on_shift(e, c, p))
            cv.bind("<Button-3>", lambda e, c=cv, p=pane: self._on_right(e, c, p))
            for seq in ("<Control-c>", "<Control-C>"):
                cv.bind(seq, lambda e: self.copy())
            for seq in ("<Control-a>", "<Control-A>"):
                cv.bind(seq, lambda e: self.select_all())
            cv.bind("<Control-Right>", lambda e: self._emit_move("AB"))
            cv.bind("<Control-Left>", lambda e: self._emit_move("BA"))
            cv.bind("<Escape>", lambda e: self.clear_selection())

    def _wheel(self, event, src):
        src.yview_scroll(-1 * (event.delta // 120), "units")
        self._sync_v_from(src)

    def _hwheel(self, event, src):
        cv, hdr = (self.cvL, self.hdrL) if src is self.cvL else (self.cvR, self.hdrR)
        cv.xview_scroll(-1 * (event.delta // 120), "units")
        hdr.xview_moveto(cv.xview()[0])

    def _vscroll(self, *args):
        self.cvL.yview(*args)
        self._sync_v_from(self.cvL)

    def _sync_v_from(self, src):
        f0, f1 = src.yview()
        for c in (self.cvL, self.cvR, self.gutter):
            if c is not src:
                c.yview_moveto(f0)
        self.vsb.set(f0, f1)

    def _hscrollL(self, *args):
        self.cvL.xview(*args)
        self.hdrL.xview(*args)

    def _hscrollR(self, *args):
        self.cvR.xview(*args)
        self.hdrR.xview(*args)

    # ---- 数据设置 ----
    def set_data(self, aligned, cols):
        self.aligned = aligned
        self.cols = cols
        self.sel = set()
        self.sel_anchor = None
        self._compute_widths()
        self._draw()

    def clear(self):
        self.aligned = []
        self.cols = []
        self.sel = set()
        self.sel_anchor = None
        for c in (self.cvL, self.cvR, self.gutter, self.hdrL, self.hdrR):
            c.delete("all")
            c.configure(scrollregion=(0, 0, 1, 1))

    def _compute_widths(self):
        self.colw = []
        for j, col in enumerate(self.cols):
            maxlen = len(_to_str(col))
            for e in self.aligned:
                v = None
                if e["a"] is not None and col in e["a"]:
                    v = e["a"].get(col, "")
                elif e["b"] is not None and col in e["b"]:
                    v = e["b"].get(col, "")
                if v is not None:
                    maxlen = max(maxlen, len(_to_str(v)))
            w = max(60, min(280, maxlen * 7 + 12))
            self.colw.append(w)

    def _clip(self, s, maxw):
        s = _to_str(s)
        maxchars = max(1, int((maxw - 8) // 7))
        if len(s) > maxchars:
            return s[:maxchars - 1] + "…"
        return s

    def _draw(self):
        self.cvL.delete("all")
        self.cvR.delete("all")
        self.gutter.delete("all")
        self.hdrL.delete("all")
        self.hdrR.delete("all")

        self.sel = set()
        self.sel_anchor = None
        total_w = self.idx_w + sum(self.colw)
        self.total_w = total_w
        total_h = self.hdrh + len(self.aligned) * self.rowh
        for cv in (self.cvL, self.cvR, self.hdrL, self.hdrR):
            cv.configure(scrollregion=(0, 0, total_w, total_h))
        self.gutter.configure(scrollregion=(0, 0, self.gw, total_h))

        self._draw_header(self.hdrL)
        self._draw_header(self.hdrR)

        sym_map = {"changed": "≠", "removed": "<", "added": ">", "same": ""}
        for i, e in enumerate(self.aligned):
            y = self.hdrh + i * self.rowh
            self._draw_row(i, e, y)
            self.cvL.create_line(0, y + self.rowh, total_w, y + self.rowh, fill="#eee")
            self.cvR.create_line(0, y + self.rowh, total_w, y + self.rowh, fill="#eee")
            sym = sym_map[e["status"]]
            if sym:
                self.gutter.create_text(self.gw / 2, y + self.rowh / 2,
                                        text=sym, fill="#c00", font=self.font)
            else:
                self.gutter.create_line(0, y + self.rowh, self.gw, y + self.rowh, fill="#f0f0f0")

        # 复位滚动
        for c in (self.cvL, self.cvR, self.gutter, self.hdrL, self.hdrR):
            c.xview_moveto(0)
            c.yview_moveto(0)
        self.vsb.set(0, 1)
        self._redraw_sel()

    def _draw_header(self, canvas):
        canvas.create_rectangle(0, 0, self.idx_w, self.hdrh, fill="#d8dee8", outline="#c8c8c8")
        canvas.create_text(4, self.hdrh / 2, text="#", anchor="w", font=self.font)
        x = self.idx_w
        for j, col in enumerate(self.cols):
            w = self.colw[j]
            canvas.create_rectangle(x, 0, x + w, self.hdrh, fill="#d8dee8", outline="#c8c8c8")
            canvas.create_text(x + 4, self.hdrh / 2, text=self._clip(col, w),
                               anchor="w", font=self.font)
            x += w

    def _draw_row(self, i, e, y):
        st = e["status"]
        changed = set(e["changed"])
        a, b = e["a"], e["b"]
        # 行号列
        self._cell(self.cvL, 0, y, self.idx_w, str(i + 1), "#f5f7fa", "#888")
        self._cell(self.cvR, 0, y, self.idx_w, str(i + 1), "#f5f7fa", "#888")
        x = self.idx_w
        for j, col in enumerate(self.cols):
            w = self.colw[j]
            va = _to_str(a.get(col, "")) if a else ""
            vb = _to_str(b.get(col, "")) if b else ""
            la_bg, la_fg = "#ffffff", "#000000"
            ra_bg, ra_fg = "#ffffff", "#000000"
            if st == "changed" and col in changed:
                la_bg, la_fg = "#ffd6d6", "#c00"
                ra_bg, ra_fg = "#ffd6d6", "#c00"
            elif st == "removed":
                la_bg, la_fg = "#ffe0e0", "#c00"
                ra_bg = "#f3f3f3"
                vb = ""
            elif st == "added":
                ra_bg, ra_fg = "#ffe0e0", "#c00"
                la_bg = "#f3f3f3"
                va = ""
            self._cell(self.cvL, x, y, w, self._clip(va, w), la_bg, la_fg)
            self._cell(self.cvR, x, y, w, self._clip(vb, w), ra_bg, ra_fg)
            x += w

    def _cell(self, cv, x, y, w, text, bg, fg):
        cv.create_rectangle(x, y, x + w, y + self.rowh, fill=bg, outline="#eee")
        cv.create_text(x + 4, y + self.rowh / 2, text=text, anchor="w",
                       fill=fg, font=self.font)

    # ---- 选中 ----
    def _build_menu(self):
        self.menu = tk.Menu(self, tearoff=0)
        self.menu.add_command(label="复制选中行（当前栏）\tCtrl+C", command=self.copy)
        self.menu.add_command(label="复制左栏 A（基准）", command=lambda: self.copy("L"))
        self.menu.add_command(label="复制右栏 B（对比）", command=lambda: self.copy("R"))
        self.menu.add_command(label="复制（含表头）", command=lambda: self.copy(header=True))
        self.menu.add_command(label="复制为差异文本", command=self.copy_diff_text)
        self.menu.add_separator()
        self.menu.add_command(label="全选\tCtrl+A", command=self.select_all)
        self.menu.add_command(label="清除选择\tEsc", command=self.clear_selection)
        self.menu.add_separator()
        self.menu.add_command(label="▶ 用左栏 A 覆盖/追加到右栏 B（改文件B）  Ctrl+→",
                              command=lambda: self._emit_move("AB"))
        self.mi_ab = self.menu.index("end")
        self.menu.add_command(label="◀ 用右栏 B 覆盖/追加到左栏 A（改文件A）  Ctrl+←",
                              command=lambda: self._emit_move("BA"))
        self.mi_ba = self.menu.index("end")
        self.menu.add_command(label="✕ 从左栏 A 删除选中行（改文件A）",
                              command=lambda: self._emit_move("delA"))
        self.mi_dela = self.menu.index("end")
        self.menu.add_command(label="✕ 从右栏 B 删除选中行（改文件B）",
                              command=lambda: self._emit_move("delB"))
        self.mi_delb = self.menu.index("end")

    def _emit_move(self, direction):
        if self.on_move is None:
            self._status("当前界面未接写回功能")
            return
        if not self.sel:
            self._status("请先选中行")
            return
        self.on_move(direction, sorted(self.sel))

    def _update_move_states(self):
        """按当前选中的行，灰化用不上的写回菜单项。"""
        rows = [self.aligned[i] for i in self.sel if 0 <= i < len(self.aligned)]
        has_a = any(r["a"] is not None for r in rows)
        has_b = any(r["b"] is not None for r in rows)
        for mi, ok in ((self.mi_ab, has_a), (self.mi_ba, has_b),
                       (self.mi_dela, has_a), (self.mi_delb, has_b)):
            self.menu.entryconfigure(mi, state="normal" if ok else "disabled")

    def _row_at(self, cv, event):
        y = cv.canvasy(event.y)
        i = int((y - self.hdrh) // self.rowh)
        if 0 <= i < len(self.aligned):
            return i
        return None

    def _on_press(self, event, cv, pane):
        cv.focus_set()
        i = self._row_at(cv, event)
        if i is None:
            return
        self.sel_pane = pane
        self.sel_anchor = i
        self._dragging = True
        self.sel = {i}
        self._redraw_sel()

    def _on_motion(self, event, cv):
        if not self._dragging or self.sel_anchor is None:
            return
        i = self._row_at(cv, event)
        if i is None:
            return
        lo, hi = sorted((self.sel_anchor, i))
        self.sel = set(range(lo, hi + 1))
        self._redraw_sel()

    def _on_release(self, event, cv):
        self._dragging = False

    def _on_ctrl(self, event, cv, pane):
        cv.focus_set()
        i = self._row_at(cv, event)
        if i is None:
            return
        self.sel_pane = pane
        self.sel_anchor = i
        self._dragging = False
        if i in self.sel:
            self.sel.discard(i)
        else:
            self.sel.add(i)
        self._redraw_sel()

    def _on_shift(self, event, cv, pane):
        cv.focus_set()
        i = self._row_at(cv, event)
        if i is None:
            return
        self.sel_pane = pane
        self._dragging = False
        if self.sel_anchor is None:
            self.sel_anchor = i
            self.sel = {i}
        else:
            lo, hi = sorted((self.sel_anchor, i))
            self.sel = set(range(lo, hi + 1))
        self._redraw_sel()

    def _on_right(self, event, cv, pane):
        i = self._row_at(cv, event)
        if i is not None and i not in self.sel:
            self.sel_pane = pane
            self.sel_anchor = i
            self.sel = {i}
            self._redraw_sel()
        if self.sel:
            self._update_move_states()
            try:
                self.menu.tk_popup(event.x_root, event.y_root)
            finally:
                self.menu.grab_release()

    def select_all(self):
        self.sel = set(range(len(self.aligned)))
        self._redraw_sel()

    def clear_selection(self):
        self.sel = set()
        self.sel_anchor = None
        self._redraw_sel()

    def _redraw_sel(self):
        for cv in (self.cvL, self.cvR):
            cv.delete("sel")
        if not self.sel:
            return
        w = max(self.total_w, self.cvL.winfo_width() or 0, 1)
        for i in self.sel:
            y = self.hdrh + i * self.rowh
            for cv in (self.cvL, self.cvR):
                cv.create_rectangle(0, y, w, y + self.rowh, fill="#1e6feb",
                                    stipple="gray25", outline="", tags="sel")
                cv.create_rectangle(0, y, w, y + self.rowh, outline="#1e6feb",
                                    width=2, tags="sel")

    # ---- 复制 ----
    def _tsv(self, d):
        if not d:
            return ""
        return "\t".join(_to_str(d.get(c, "")) for c in self.cols)

    def copy(self, side=None, header=False):
        if not self.sel:
            self._status("请先选中行再复制")
            return
        side = side or self.sel_pane
        lines = []
        if header:
            lines.append("\t".join(self.cols))
        for i in sorted(self.sel):
            e = self.aligned[i]
            if side == "L":
                d = e["a"] if e["a"] is not None else e["b"]
            else:
                d = e["b"] if e["b"] is not None else e["a"]
            lines.append(self._tsv(d))
        self._to_clipboard("\n".join(lines), len(self.sel))

    def copy_diff_text(self):
        if not self.sel:
            self._status("请先选中行再复制")
            return
        lines = []
        for i in sorted(self.sel):
            e = self.aligned[i]
            st = e["status"]
            if st == "changed":
                for c in e["changed"]:
                    lines.append("修改 | %s | %s: %s → %s" % (
                        e["key"], c, _to_str(e["a"].get(c, "")), _to_str(e["b"].get(c, ""))))
            elif st == "added":
                lines.append("新增 | %s | %s" % (e["key"], self._tsv(e["b"])))
            elif st == "removed":
                lines.append("删除 | %s | %s" % (e["key"], self._tsv(e["a"])))
        if not lines:
            self._status("所选行没有差异")
            return
        self._to_clipboard("\n".join(lines), len(lines))

    def _to_clipboard(self, text, n):
        self.clipboard_clear()
        self.clipboard_append(text)
        try:
            self.update_idletasks()
        except Exception:
            pass
        self._status("已复制 %d 行到剪贴板" % n)

    def _status(self, msg):
        if callable(self.on_status):
            self.on_status(msg)


# ----------------------------------------------------------------------------
# 主窗口
# ----------------------------------------------------------------------------

class App:
    def __init__(self, root):
        self.root = root
        self.root.title("ExcelDiff - 两个 Excel 内容对比")
        self.root.geometry("1180x720")
        try:
            self.root.iconbitmap(self._resource("icon.ico"))
        except Exception:
            pass

        self.file_a = tk.StringVar()
        self.file_b = tk.StringVar()
        self.sheet_a = tk.StringVar()
        self.sheet_b = tk.StringVar()
        self.mode = tk.StringVar(value="index")
        self.key_col = tk.StringVar()
        self.header_row = tk.IntVar(value=1)
        self.ignore_case = tk.BooleanVar(value=False)
        self.skip_empty = tk.BooleanVar(value=True)
        self.last_dir = os.path.expanduser("~")

        self._build_ui()
        self.result = None
        self._last_files = ("", "", "", "")
        self._backups = {}        # 源文件 -> 本次会话的备份路径
        self._ui_queue = queue.Queue()
        self.root.after(60, self._drain_ui_queue)

    # ---- 工作线程 -> 主线程 的回调队列（Tk 只能在主线程操作）----
    def _post(self, fn):
        self._ui_queue.put(fn)

    def _drain_ui_queue(self):
        try:
            while True:
                fn = self._ui_queue.get_nowait()
                try:
                    fn()
                except Exception:
                    pass
        except queue.Empty:
            pass
        try:
            self.root.after(60, self._drain_ui_queue)
        except Exception:
            pass

    def _resource(self, name):
        if getattr(sys, "frozen", False):
            return os.path.join(os.path.dirname(sys.executable), name)
        return os.path.join(os.path.dirname(os.path.abspath(__file__)), name)

    def _build_ui(self):
        pad = {"padx": 8, "pady": 4}

        frm_top = ttk.LabelFrame(self.root, text="文件选择", padding=8)
        frm_top.pack(fill="x", **pad)

        ttk.Label(frm_top, text="文件A（基准）").grid(row=0, column=0, sticky="w")
        ttk.Entry(frm_top, textvariable=self.file_a, width=60).grid(row=0, column=1, sticky="ew")
        ttk.Button(frm_top, text="浏览…", command=lambda: self._pick(self.file_a, self.sheet_a)).grid(row=0, column=2)
        ttk.Label(frm_top, text="工作表").grid(row=0, column=3, padx=(12, 2))
        self.cb_sheet_a = ttk.Combobox(frm_top, textvariable=self.sheet_a, width=14, state="readonly")
        self.cb_sheet_a.grid(row=0, column=4)

        ttk.Label(frm_top, text="文件B（对比）").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(frm_top, textvariable=self.file_b, width=60).grid(row=1, column=1, sticky="ew", pady=(6, 0))
        ttk.Button(frm_top, text="浏览…", command=lambda: self._pick(self.file_b, self.sheet_b)).grid(row=1, column=2, pady=(6, 0))
        ttk.Label(frm_top, text="工作表").grid(row=1, column=3, padx=(12, 2), pady=(6, 0))
        self.cb_sheet_b = ttk.Combobox(frm_top, textvariable=self.sheet_b, width=14, state="readonly")
        self.cb_sheet_b.grid(row=1, column=4, pady=(6, 0))
        frm_top.columnconfigure(1, weight=1)

        frm_opt = ttk.LabelFrame(self.root, text="比对选项", padding=8)
        frm_opt.pack(fill="x", **pad)

        ttk.Radiobutton(frm_opt, text="按行号对齐", variable=self.mode,
                        value="index", command=self._on_mode).grid(row=0, column=0)
        ttk.Radiobutton(frm_opt, text="按关键字列", variable=self.mode,
                        value="key", command=self._on_mode).grid(row=0, column=1)
        ttk.Label(frm_opt, text="关键字列").grid(row=0, column=2, padx=(12, 2))
        self.cb_key = ttk.Combobox(frm_opt, textvariable=self.key_col, width=16, state="readonly")
        self.cb_key.grid(row=0, column=3)
        ttk.Label(frm_opt, text="表头行").grid(row=0, column=4, padx=(12, 2))
        ttk.Spinbox(frm_opt, from_=1, to=20, width=5, textvariable=self.header_row).grid(row=0, column=5)
        ttk.Checkbutton(frm_opt, text="忽略大小写", variable=self.ignore_case).grid(row=0, column=6, padx=(12, 0))
        ttk.Checkbutton(frm_opt, text="跳过空行", variable=self.skip_empty).grid(row=0, column=7, padx=(8, 0))

        frm_btn = ttk.Frame(self.root, padding=(8, 0))
        frm_btn.pack(fill="x")
        ttk.Button(frm_btn, text="开始比对", command=self._start).pack(side="left")
        ttk.Button(frm_btn, text="导出差异 Excel", command=self._export_xlsx).pack(side="left", padx=6)
        ttk.Button(frm_btn, text="生成 HTML 报告", command=self._export_html).pack(side="left")
        ttk.Button(frm_btn, text="清空", command=self._clear).pack(side="left", padx=6)
        ttk.Separator(frm_btn, orient="vertical").pack(side="left", fill="y", padx=8, pady=2)
        ttk.Label(frm_btn, text="行写回：").pack(side="left")
        self.btn_ab = ttk.Button(frm_btn, text="A → B  写入B",
                                 command=lambda: self._apply_moves("AB"))
        self.btn_ab.pack(side="left", padx=(4, 4))
        self.btn_ba = ttk.Button(frm_btn, text="B → A  写入A",
                                 command=lambda: self._apply_moves("BA"))
        self.btn_ba.pack(side="left")

        self.status = tk.StringVar(value="请选择两个 Excel 文件后点击“开始比对”")
        ttk.Label(self.root, textvariable=self.status, relief="sunken", anchor="w").pack(fill="x", side="bottom")

        frm_res = ttk.Frame(self.root, padding=8)
        frm_res.pack(fill="both", expand=True)
        self.compare_view = CompareView(frm_res)
        self.compare_view.on_status = self.status.set
        self.compare_view.on_move = self._apply_moves
        self.compare_view.pack(fill="both", expand=True)

        self._on_mode()

    def _on_mode(self):
        if self.mode.get() == "key":
            self.cb_key.configure(state="readonly")
            if self.file_a.get() and os.path.exists(self.file_a.get()):
                try:
                    h, _ = load_excel(self.file_a.get(),
                                      self.sheet_a.get() or None, self.header_row.get())
                    self.cb_key["values"] = [x for x in h if x != ""]
                    if not self.key_col.get() and self.cb_key["values"]:
                        self.key_col.set(self.cb_key["values"][0])
                except Exception:
                    pass
        else:
            self.cb_key.set("")
            self.cb_key.configure(state="disabled")

    def _pick(self, var, sheet_var):
        path = filedialog.askopenfilename(
            title="选择 Excel 文件",
            initialdir=self.last_dir,
            filetypes=[("Excel 文件", "*.xlsx *.xlsm *.xls"), ("所有文件", "*.*")])
        if not path:
            return
        self.last_dir = os.path.dirname(path)
        var.set(path)
        try:
            names = self._sheet_names(path)
            cb = self.cb_sheet_a if var is self.file_a else self.cb_sheet_b
            cb["values"] = names
            if names:
                sheet_var.set(names[0])
        except Exception as e:
            messagebox.showerror("读取失败", str(e))
        self._on_mode()

    def _sheet_names(self, path):
        ext = os.path.splitext(path)[1].lower()
        if ext == ".xls":
            book = xlrd.open_workbook(path)
            return book.sheet_names()
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            return wb.sheetnames
        finally:
            wb.close()

    def _start(self):
        fa, fb = self.file_a.get(), self.file_b.get()
        if not fa or not fb:
            messagebox.showwarning("提示", "请先选择两个 Excel 文件")
            return
        if not os.path.exists(fa) or not os.path.exists(fb):
            messagebox.showerror("错误", "文件路径不存在，请重新选择")
            return
        if self.mode.get() == "key" and not self.key_col.get():
            messagebox.showwarning("提示", "请选择关键字列")
            return

        self.status.set("正在比对…")
        self.root.update_idletasks()

        # Tk 变量只能在主线程读，先取值再进线程
        sa, sb = self.sheet_a.get(), self.sheet_b.get()
        hr = self.header_row.get()
        mode, key = self.mode.get(), self.key_col.get()
        ic, se = self.ignore_case.get(), self.skip_empty.get()

        def work():
            try:
                ha, da, ra = load_excel_ex(fa, sa or None, hr)
                hb, db, rb = load_excel_ex(fb, sb or None, hr)
                res = compare(ha, da, hb, db, mode, key, ic, se, ra, rb)
                self.result = res
                self._last_files = (fa, sa, fb, sb)
                self._post(lambda: self._show_result(fa, fb, res))
            except Exception as e:
                err = str(e)
                self._post(lambda: messagebox.showerror("比对失败", err))
                self._post(lambda: self.status.set("比对失败：" + err))

        threading.Thread(target=work, daemon=True).start()

    def _show_result(self, fa, fb, res):
        self.compare_view.set_data(res["aligned"], res["unified"])
        s = res["stats"]
        total = len(res["aligned"])
        msg = ("比对完成 ｜ 共 %d 行 ｜ 相同 %d ｜ 修改 %d ｜ 新增 %d ｜ 删除 %d ｜ 差异合计 %d"
               % (total, s["same"], s["changed"], s["added"], s["removed"],
                  s["changed"] + s["added"] + s["removed"]))
        if res["warnings"]:
            msg += "  ｜ 注意：" + "；".join(res["warnings"])
        self.status.set(msg)

    def _opts_text(self):
        if self.mode.get() == "index":
            return "按行号对齐（表头行=%d）" % self.header_row.get()
        return "按关键字列「%s」（表头行=%d）" % (self.key_col.get(), self.header_row.get())

    # ---- 行写回（把行"移到"另一侧，直接改源文件）----
    def _sel_rows(self):
        if not self.result:
            return []
        aligned = self.result["aligned"]
        return [aligned[i] for i in sorted(self.compare_view.sel)
                if 0 <= i < len(aligned)]

    def _apply_moves(self, direction):
        if not self.result:
            messagebox.showwarning("提示", "请先执行比对")
            return
        rows = self._sel_rows()
        if not rows:
            messagebox.showwarning("提示",
                                   "请先在对比表里选中要操作的行\n（单击选中，Ctrl 加选，Shift 连选）")
            return

        fa, sa, fb, sb = self._last_files
        hr = self.header_row.get()
        to_b = direction in ("AB", "delB")
        tpath, tsheet, tname = (fb, sb, "B") if to_b else (fa, sa, "A")
        if not tpath:
            return
        if os.path.splitext(tpath)[1].lower() == ".xls":
            messagebox.showerror("无法写回",
                                 "文件%s 是旧版 .xls，程序不能写入。\n"
                                 "请先用 Excel 另存为 .xlsx 后重新比对。" % tname)
            return

        unified = self.result["unified"]
        edits, skipped = [], []
        for e in rows:
            if direction in ("delA", "delB"):
                rn = e["ra"] if direction == "delA" else e["rb"]
                if rn:
                    edits.append({"op": "delete", "row": rn})
                else:
                    skipped.append(e["key"])
                continue
            src = e["a"] if to_b else e["b"]
            has_tgt = (e["b"] is not None) if to_b else (e["a"] is not None)
            if e["status"] == "same":
                # 两侧本来就一致：不写，避免把目标侧的公式覆盖成静态值
                skipped.append(e["key"])
                continue
            if src is None:
                skipped.append(e["key"])
                continue
            if not has_tgt:
                # 该行只有一侧有 —— 追加到目标文件表尾
                edits.append({"op": "insert", "row": None,
                              "values": {c: src.get(c, "") for c in unified}})
            else:
                rn = e["rb"] if to_b else e["ra"]
                if not rn:
                    skipped.append(e["key"])
                    continue
                cols = list(e["changed"]) if e["status"] == "changed" else list(unified)
                edits.append({"op": "set", "row": rn,
                              "values": {c: src.get(c, "") for c in cols}})

        if not edits:
            messagebox.showinfo("无需操作", "选中的行没有可写入的内容。")
            return

        n_set = sum(1 for x in edits if x["op"] == "set")
        n_ins = sum(1 for x in edits if x["op"] == "insert")
        n_del = sum(1 for x in edits if x["op"] == "delete")
        parts = []
        if n_set:
            parts.append("覆盖 %d 行" % n_set)
        if n_ins:
            parts.append("追加 %d 行到表尾" % n_ins)
        if n_del:
            parts.append("删除 %d 行（不可撤销，只能从备份恢复）" % n_del)
        src_desc = {"AB": "用左栏 A 的值写入文件B",
                    "BA": "用右栏 B 的值写入文件A",
                    "delA": "从文件A 删除这些行",
                    "delB": "从文件B 删除这些行"}[direction]
        msg = ("即将直接修改源文件：\n\n"
               "文件%s：%s\n工作表：%s\n\n%s\n\n"
               "写入前会自动备份为「原文件名.bak-时间戳」。\n是否继续？"
               % (tname, tpath, tsheet or "(默认)", src_desc + "；" + "、".join(parts)))
        if skipped:
            msg += "\n\n（跳过 %d 行：该侧没有对应记录或取不到行号）" % len(skipped)
        if not messagebox.askyesno("确认写回源文件", msg):
            return

        self.status.set("正在写入文件%s…" % tname)
        self.root.update_idletasks()
        keys_written = {e["key"] for e in rows}
        mode, key_col = self.mode.get(), self.key_col.get()
        ic, se = self.ignore_case.get(), self.skip_empty.get()

        def work():
            try:
                if tpath not in self._backups:
                    self._backups[tpath] = backup_file(tpath)
                done = apply_edits(tpath, tsheet or None, hr, edits)
                ha, da, ra = load_excel_ex(fa, sa or None, hr)
                hb, db, rb = load_excel_ex(fb, sb or None, hr)
                res = compare(ha, da, hb, db, mode, key_col, ic, se, ra, rb)
                self.result = res
                self._post(lambda: self._after_write(res, tname, len(done),
                                                     keys_written, tpath))
            except Exception as ex:
                err = str(ex)
                self._post(lambda: self._write_failed(err))

        threading.Thread(target=work, daemon=True).start()

    def _after_write(self, res, tname, n, keys_written, tpath):
        self.compare_view.set_data(res["aligned"], res["unified"])
        idx = {e["key"]: i for i, e in enumerate(res["aligned"])}
        self.compare_view.sel = {idx[k] for k in keys_written if k in idx}
        self.compare_view._redraw_sel()
        s = res["stats"]
        bak = self._backups.get(tpath, "")
        self.status.set("已写入文件%s（%d 项）｜ 剩余差异：修改 %d ｜ 新增 %d ｜ 删除 %d ｜ 备份：%s"
                        % (tname, n, s["changed"], s["added"], s["removed"], bak))

    def _write_failed(self, err):
        low = err.lower()
        if "permission" in low or "being used" in low or "另一个程序" in err or "拒绝访问" in err:
            messagebox.showerror("写回失败",
                                 "文件可能正在 Excel 中打开，请关闭该文件后重试。\n\n" + err)
        else:
            messagebox.showerror("写回失败", err)
        self.status.set("写回失败：" + err)

    def _export_xlsx(self):
        if not self.result:
            messagebox.showwarning("提示", "请先执行比对")
            return
        path = filedialog.asksaveasfilename(
            title="导出差异 Excel", initialdir=self.last_dir,
            defaultextension=".xlsx", filetypes=[("Excel 工作簿", "*.xlsx")])
        if not path:
            return
        try:
            fa, sa, fb, sb = self._last_files
            info = {"file_a": fa or "", "sheet_a": sa or "(默认)",
                    "file_b": fb or "", "sheet_b": sb or "(默认)",
                    "opts": self._opts_text()}
            export_excel(self.result, path, info)
            messagebox.showinfo("完成", "已导出：\n" + path)
        except Exception as e:
            messagebox.showerror("导出失败", str(e))

    def _export_html(self):
        if not self.result:
            messagebox.showwarning("提示", "请先执行比对")
            return
        fa, sa, fb, sb = self._last_files
        path = filedialog.asksaveasfilename(
            title="生成 HTML 报告", initialdir=self.last_dir,
            defaultextension=".html", filetypes=[("HTML 文件", "*.html")])
        if not path:
            return
        try:
            html = build_html(self.result, fa, sa or "(默认)", fb, sb or "(默认)", self._opts_text())
            with open(path, "w", encoding="utf-8") as f:
                f.write(html)
            webbrowser.open(path)
            messagebox.showinfo("完成", "报告已生成并在浏览器打开：\n" + path)
        except Exception as e:
            messagebox.showerror("生成失败", str(e))

    def _clear(self):
        self.compare_view.clear()
        self.result = None
        self.status.set("已清空结果")


def _selftest(out_path):
    """内置自检：生成两个临时 xlsx 并跑一次比对，结果写入文件。用于验证打包后的运行时。"""
    import tempfile

    def mk(p, rows, h):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(h)
        for r in rows:
            ws.append(r)
        wb.save(p)

    d = tempfile.mkdtemp()
    a = os.path.join(d, "a.xlsx")
    b = os.path.join(d, "b.xlsx")
    mk(a, [["K1", 1], ["K2", 2]], ["id", "v"])
    mk(b, [["K1", 1], ["K2", 9], ["K3", 3]], ["id", "v"])
    ha, da, ra = load_excel_ex(a, None, 1)
    hb, db, rb = load_excel_ex(b, None, 1)
    res = compare(ha, da, hb, db, "key", "id", False, True, ra, rb)
    s = res["stats"]
    ok = (s["same"] == 1 and s["changed"] == 1 and s["added"] == 1 and s["removed"] == 0)

    # 同时验证导出 + 写回通道（含样式写入），覆盖打包后运行时
    xp = os.path.join(d, "out.xlsx")
    hp = os.path.join(d, "out.html")
    export_excel(res, xp, {"file_a": a, "file_b": b, "opts": "selftest"})
    with open(hp, "w", encoding="utf-8") as f:
        f.write(build_html(res, a, "Sheet", b, "Sheet", "selftest"))
    hot = 0
    try:
        wb = load_workbook(xp)
        ws = wb["差异明细"]
        for row in ws.iter_rows():
            for c in row:
                if c.fill and c.fill.fill_type == "solid" and c.fill.fgColor.rgb == "00FDE2E1":
                    hot += 1
        wb.close()
    except Exception as e:
        ok = False
        hot = "ERR:%s" % e
    if not (os.path.getsize(xp) > 0 and os.path.getsize(hp) > 0):
        ok = False

    # 写回通道：在临时副本上做一次覆盖，验证冻结环境能保存 xlsx
    wr = "skip"
    try:
        c = os.path.join(d, "c.xlsx")
        shutil.copy2(b, c)
        hc, dc, rc = load_excel_ex(c, None, 1)
        e0 = [x for x in res["aligned"] if x["status"] == "changed"][0]
        apply_edits(c, None, 1, [{"op": "set", "row": e0["rb"],
                                  "values": {k: e0["a"].get(k) for k in e0["changed"]}}])
        hc2, dc2, _ = load_excel_ex(c, None, 1)
        r2 = compare(hc2, dc2, ha, da, "key", "id", False, True)
        wr = "ok" if r2["stats"]["changed"] == 0 else "FAIL %s" % r2["stats"]
        if wr != "ok":
            ok = False
    except Exception as e:
        ok = False
        wr = "ERR:%s" % e

    with open(out_path, "w", encoding="utf-8") as f:
        f.write(("PASS" if ok else "FAIL") + " stats=%s aligned=%d hot_cells=%s write=%s\n"
                % (s, len(res["aligned"]), hot, wr))


def _read_probe(path, out_path):
    """诊断模式：读取指定 Excel 并报告各工作表行数，结果写入文件。"""
    lines = []
    try:
        ext = os.path.splitext(path)[1].lower()
        if ext == ".xls":
            names = xlrd.open_workbook(path).sheet_names()
        else:
            wb = load_workbook(path, read_only=True, data_only=True)
            names = wb.sheetnames
            wb.close()
        lines.append("file=%s" % path)
        lines.append("sheets=%s" % names)
        for sn in names:
            try:
                h, d = load_excel(path, sn, 1)
                lines.append("  [%s] cols=%d data_rows=%d header=%s"
                             % (sn, len(h), len(d), h[:8]))
            except Exception as e:
                lines.append("  [%s] ERROR: %s" % (sn, e))
    except Exception as e:
        lines.append("FATAL: %s" % e)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")


def main():
    if "--read" in sys.argv:
        i = sys.argv.index("--read")
        path = sys.argv[i + 1]
        out = sys.argv[i + 2] if i + 2 < len(sys.argv) else "read_out.txt"
        _read_probe(path, out)
        return
    if "--selftest" in sys.argv:
        i = sys.argv.index("--selftest")
        out = sys.argv[i + 1] if i + 1 < len(sys.argv) else "selftest_out.txt"
        _selftest(out)
        return
    root = tk.Tk()
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
